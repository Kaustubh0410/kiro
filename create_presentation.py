"""
Create Smart Food Management PowerPoint Presentation
Generates a 10-slide presentation with JTP template format and embedded images.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# =============================================================================
# Global Constants
# =============================================================================
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Image directory
IMAGES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "images")

# Colors
DARK_BG = RGBColor(0x1B, 0x28, 0x38)
CARD_BG = RGBColor(0x2C, 0x3E, 0x50)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
ORANGE = RGBColor(0xF3, 0x9C, 0x12)
RED = RGBColor(0xE7, 0x4C, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xEC, 0xF0, 0xF1)
FOOTER_GRAY = RGBColor(0x95, 0xA5, 0xA6)
BLUE = RGBColor(0x34, 0x98, 0xDB)
PURPLE = RGBColor(0x9B, 0x59, 0xB6)
OVERLAY_COLOR = RGBColor(0x1B, 0x28, 0x38)

FOOTER_TEXT = "(c) JTP Co., Ltd. All Rights Reserved."


def add_dark_background(slide):
    """Add a full-slide dark background rectangle."""
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    return bg


def add_background_image(slide, image_filename):
    """Add an image as a full-slide background."""
    image_path = os.path.join(IMAGES_DIR, image_filename)
    if os.path.exists(image_path):
        pic = slide.shapes.add_picture(
            image_path, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT
        )
        return pic
    print(f"WARNING: Image not found: {image_path}")
    return None


def add_dark_overlay(slide, opacity_level="heavy"):
    """Add a dark overlay on top of background image for text readability."""
    overlay = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT
    )
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = OVERLAY_COLOR
    overlay.line.fill.background()
    # Set transparency via the XML element
    # Access the spPr/solidFill element in the shape XML
    sp_elem = overlay._element
    sp_pr = sp_elem.find(qn('p:spPr'))
    if sp_pr is None:
        sp_pr = sp_elem.find('.//' + qn('p:spPr'))
    solid_fill = sp_pr.find(qn('a:solidFill'))
    if solid_fill is not None:
        srgb = solid_fill.find(qn('a:srgbClr'))
        if srgb is not None:
            # Remove existing alpha if any
            for existing_alpha in srgb.findall(qn('a:alpha')):
                srgb.remove(existing_alpha)
            alpha_elem = etree.SubElement(srgb, qn('a:alpha'))
            if opacity_level == "heavy":
                alpha_elem.set('val', '75000')  # 75% opacity
            elif opacity_level == "medium":
                alpha_elem.set('val', '60000')  # 60% opacity
            else:
                alpha_elem.set('val', '50000')  # 50% opacity
    return overlay


def add_footer(slide):
    """Add the JTP copyright footer centered at the bottom."""
    footer = slide.shapes.add_textbox(
        Inches(0), Inches(7.0), SLIDE_WIDTH, Inches(0.4)
    )
    tf = footer.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = FOOTER_TEXT
    run.font.size = Pt(10)
    run.font.color.rgb = FOOTER_GRAY


def add_section_label(slide, text):
    """Add the UPPERCASE section label at top-left in small bold accent text."""
    label_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.4))
    tf = label_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text.upper()
    run.font.size = Pt(11)
    run.font.color.rgb = ORANGE
    run.font.bold = True


def set_text_props(text_frame, text, font_size, color, bold=False, alignment=PP_ALIGN.LEFT):
    """Helper to set text properties on a text frame's first paragraph."""
    tf = text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    return tf


def add_paragraph(text_frame, text, font_size, color, bold=False, alignment=PP_ALIGN.LEFT):
    """Add a new paragraph to an existing text frame."""
    p = text_frame.add_paragraph()
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    return p


def add_card(slide, left, top, width, height, fill_color=None, accent_color=None, accent_position="top"):
    """Add a rounded rectangle card with optional colored accent."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color or CARD_BG
    card.line.fill.background()

    if accent_color:
        if accent_position == "top":
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.08)
            )
        elif accent_position == "left":
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height
            )
        accent.fill.solid()
        accent.fill.fore_color.rgb = accent_color
        accent.line.fill.background()

    return card


def create_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1 - Title
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "01_title.jpg")
    add_dark_overlay(slide, "heavy")

    # JTP branding at top
    jtp_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.4))
    set_text_props(jtp_box.text_frame, "JTP Co., Ltd.", 14, WHITE, bold=True)

    # Author below company name
    author_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(4), Inches(0.4))
    set_text_props(author_box.text_frame, "Kaustubh Londhe", 12, LIGHT_GRAY)

    # Green accent bar
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.4), Inches(2.0), Inches(0.08)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = GREEN
    accent.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(8), Inches(1.2))
    set_text_props(title_box.text_frame, "Smart Food Management", 48, WHITE, bold=True)

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(8), Inches(0.8))
    set_text_props(sub_box.text_frame, "Intelligent Platform for Zero Food Waste", 24, LIGHT_GRAY)

    # Stat callout
    stat_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(9), Inches(0.8))
    set_text_props(stat_box.text_frame, "1.3 Billion Tonnes of food wasted annually", 28, ORANGE, bold=True)

    add_footer(slide)

    # =========================================================================
    # SLIDE 2 - The Crisis
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "02_crisis.jpg")
    add_dark_overlay(slide, "heavy")

    add_section_label(slide, "THE CRISIS")

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12), Inches(1.0))
    set_text_props(title_box.text_frame, "Every Day, Food Goes to Waste", 36, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # 3 stat cards
    cards_data = [
        ("1.3B", "Tonnes", "Food wasted yearly", ORANGE),
        ("$1 Trillion", "", "Economic loss annually", ORANGE),
        ("828 Million", "", "People go hungry", RED),
    ]
    card_width = Inches(3.5)
    card_height = Inches(3.5)
    start_left = Inches(1.2)
    spacing = Inches(0.5)

    for i, (stat, unit, desc, accent_color) in enumerate(cards_data):
        left = start_left + i * (card_width + spacing)
        top = Inches(2.2)

        add_card(slide, left, top, card_width, card_height, accent_color=accent_color, accent_position="top")

        # Stat number
        stat_tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.6), card_width - Inches(0.6), Inches(1.2))
        set_text_props(stat_tb.text_frame, stat, 40, accent_color, bold=True, alignment=PP_ALIGN.CENTER)

        # Unit
        if unit:
            unit_tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(1.6), card_width - Inches(0.6), Inches(0.6))
            set_text_props(unit_tb.text_frame, unit, 22, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

        # Description
        desc_tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(2.4), card_width - Inches(0.6), Inches(0.8))
        set_text_props(desc_tb.text_frame, desc, 18, LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_footer(slide)

    # =========================================================================
    # SLIDES 3, 4, 5 - Problem + Solution pairs
    # =========================================================================
    problem_solution_slides = [
        {
            "image": "03_fridge.jpg",
            "problem_title": "Forgotten Food",
            "problem_desc": "Households throw away food because they forget what is in their fridge",
            "solution_title": "Smart Inventory Tracking",
            "solution_desc": "Automatic reminders before food expires. Never waste forgotten items again.",
        },
        {
            "image": "04_donation.jpg",
            "problem_title": "Wasted Surplus",
            "problem_desc": "Restaurants & stores discard surplus with no way to redistribute",
            "solution_title": "Surplus Food Marketplace",
            "solution_desc": "Connects donors with food banks & NGOs instantly. Every extra meal finds a home.",
        },
        {
            "image": "05_cooking.jpg",
            "problem_title": "Unused Ingredients",
            "problem_desc": "People do not know how to use leftover ingredients creatively",
            "solution_title": "AI Recipe Suggestions",
            "solution_desc": "Creates delicious meals from what you already have. Turn leftovers into feasts.",
        },
    ]

    for ps_data in problem_solution_slides:
        slide = prs.slides.add_slide(blank_layout)
        add_background_image(slide, ps_data["image"])
        add_dark_overlay(slide, "heavy")

        add_section_label(slide, "THE CHALLENGE & SOLUTION")

        # Problem card (left half)
        p_left = Inches(0.8)
        p_top = Inches(1.5)
        p_width = Inches(5.2)
        p_height = Inches(4.8)

        add_card(slide, p_left, p_top, p_width, p_height, accent_color=RED, accent_position="top")

        # Problem label
        pl_box = slide.shapes.add_textbox(p_left + Inches(0.5), p_top + Inches(0.4), p_width - Inches(1.0), Inches(0.5))
        set_text_props(pl_box.text_frame, "THE PROBLEM", 12, RED, bold=True, alignment=PP_ALIGN.CENTER)

        # Problem title
        pt_box = slide.shapes.add_textbox(p_left + Inches(0.5), p_top + Inches(1.2), p_width - Inches(1.0), Inches(0.7))
        set_text_props(pt_box.text_frame, ps_data["problem_title"], 26, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

        # Problem description
        pd_box = slide.shapes.add_textbox(p_left + Inches(0.5), p_top + Inches(2.2), p_width - Inches(1.0), Inches(1.5))
        set_text_props(pd_box.text_frame, ps_data["problem_desc"], 16, LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

        # Arrow between
        arrow_box = slide.shapes.add_textbox(Inches(6.1), Inches(3.5), Inches(1.0), Inches(0.8))
        set_text_props(arrow_box.text_frame, "\u2192", 36, GREEN, bold=True, alignment=PP_ALIGN.CENTER)

        # Solution card (right half)
        s_left = Inches(7.2)
        s_top = Inches(1.5)
        s_width = Inches(5.2)
        s_height = Inches(4.8)

        add_card(slide, s_left, s_top, s_width, s_height, accent_color=GREEN, accent_position="top")

        # Solution label
        sl_box = slide.shapes.add_textbox(s_left + Inches(0.5), s_top + Inches(0.4), s_width - Inches(1.0), Inches(0.5))
        set_text_props(sl_box.text_frame, "THE SOLUTION", 12, GREEN, bold=True, alignment=PP_ALIGN.CENTER)

        # Solution title
        st_box = slide.shapes.add_textbox(s_left + Inches(0.5), s_top + Inches(1.2), s_width - Inches(1.0), Inches(0.7))
        set_text_props(st_box.text_frame, ps_data["solution_title"], 26, GREEN, bold=True, alignment=PP_ALIGN.CENTER)

        # Solution description
        sd_box = slide.shapes.add_textbox(s_left + Inches(0.5), s_top + Inches(2.2), s_width - Inches(1.0), Inches(1.5))
        set_text_props(sd_box.text_frame, ps_data["solution_desc"], 16, LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

        add_footer(slide)

    # =========================================================================
    # SLIDE 6 - Platform Features
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "06_app.jpg")
    add_dark_overlay(slide, "heavy")

    add_section_label(slide, "FEATURES")

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12), Inches(1.0))
    set_text_props(title_box.text_frame, "Six Superpowers in One Platform", 36, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # 6 feature cards in 2x3 grid
    features = [
        ("Smart Expiry Alerts", "Never let food expire unnoticed", GREEN),
        ("AI Recipe Suggestions", "Cook smart with what you have", ORANGE),
        ("Community Food Sharing", "Share surplus with neighbors", GREEN),
        ("Waste Analytics Dashboard", "Track and reduce your waste", ORANGE),
        ("Donation Network", "Connect with food banks easily", GREEN),
        ("Multi-platform Access", "Use on any device, anywhere", ORANGE),
    ]

    card_w = Inches(3.6)
    card_h = Inches(2.2)
    start_x = Inches(0.9)
    start_y = Inches(2.0)
    x_gap = Inches(0.4)
    y_gap = Inches(0.4)

    for i, (name, desc, accent) in enumerate(features):
        col = i % 3
        row = i // 3
        left = start_x + col * (card_w + x_gap)
        top = start_y + row * (card_h + y_gap)

        add_card(slide, left, top, card_w, card_h, accent_color=accent, accent_position="top")

        # Feature name
        n_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.5), card_w - Inches(0.4), Inches(0.8))
        set_text_props(n_box.text_frame, name, 18, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

        # Description
        d_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(1.3), card_w - Inches(0.4), Inches(0.7))
        set_text_props(d_box.text_frame, desc, 13, LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    add_footer(slide)

    # =========================================================================
    # SLIDE 7 - Architecture (User Benefits)
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "10_dashboard.jpg")
    add_dark_overlay(slide, "heavy")

    add_section_label(slide, "ARCHITECTURE")

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12), Inches(1.0))
    set_text_props(title_box.text_frame, "How It Works For You", 36, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    layers = [
        ("Works on any device", "Seamless experience on phone, tablet, or computer", BLUE),
        ("Lightning-fast & beautiful", "Instant responses with an intuitive, clean design", ORANGE),
        ("Your data is always safe", "Bank-level security protecting your information", GREEN),
        ("Smart AI that learns your habits", "Gets better and more personalized the more you use it", PURPLE),
    ]

    bar_width = Inches(11.0)
    bar_height = Inches(1.2)
    bar_start_x = Inches(1.1)
    bar_start_y = Inches(2.0)
    bar_gap = Inches(0.3)

    for i, (title, desc, accent) in enumerate(layers):
        top = bar_start_y + i * (bar_height + bar_gap)

        add_card(slide, bar_start_x, top, bar_width, bar_height, accent_color=accent, accent_position="left")

        # Title text
        t_box = slide.shapes.add_textbox(bar_start_x + Inches(0.5), top + Inches(0.15), Inches(9.5), Inches(0.6))
        set_text_props(t_box.text_frame, title, 20, WHITE, bold=True, alignment=PP_ALIGN.LEFT)

        # Description
        d_box = slide.shapes.add_textbox(bar_start_x + Inches(0.5), top + Inches(0.65), Inches(9.5), Inches(0.5))
        set_text_props(d_box.text_frame, desc, 14, LIGHT_GRAY, alignment=PP_ALIGN.LEFT)

    add_footer(slide)

    # =========================================================================
    # SLIDE 8 - Impact
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "07_community.jpg")
    add_dark_overlay(slide, "heavy")

    add_section_label(slide, "PROJECTED IMPACT")

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12), Inches(1.0))
    set_text_props(title_box.text_frame, "The Impact We Will Create", 36, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    metrics = [
        ("50%", "Less household food waste", GREEN),
        ("10M+", "Meals redistributed", ORANGE),
        ("Zero", "Food expires forgotten", GREEN),
    ]

    metric_width = Inches(3.5)
    metric_height = Inches(3.5)
    metric_start_x = Inches(1.2)
    metric_gap = Inches(0.5)

    for i, (number, desc, color) in enumerate(metrics):
        left = metric_start_x + i * (metric_width + metric_gap)
        top = Inches(2.2)

        add_card(slide, left, top, metric_width, metric_height, accent_color=color, accent_position="top")

        # Large number
        n_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.8), metric_width - Inches(0.6), Inches(1.5))
        set_text_props(n_box.text_frame, number, 54, color, bold=True, alignment=PP_ALIGN.CENTER)

        # Description
        d_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(2.3), metric_width - Inches(0.6), Inches(0.8))
        set_text_props(d_box.text_frame, desc, 18, WHITE, alignment=PP_ALIGN.CENTER)

    add_footer(slide)

    # =========================================================================
    # SLIDE 9 - Roadmap
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "08_growth.jpg")
    add_dark_overlay(slide, "heavy")

    add_section_label(slide, "THE FUTURE")

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.8), Inches(12), Inches(1.0))
    set_text_props(title_box.text_frame, "Our Journey Forward", 36, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    phases = [
        ("1", "Smart Inventory\n& Alerts", GREEN),
        ("2", "AI Recipes &\nCommunity Sharing", ORANGE),
        ("3", "Marketplace &\nDonation Network", BLUE),
        ("4", "Global Expansion\n& Partnerships", PURPLE),
    ]

    phase_width = Inches(2.7)
    phase_height = Inches(3.2)
    phase_start_x = Inches(0.7)
    phase_gap = Inches(0.4)
    phase_top = Inches(2.2)

    # Connecting line
    line_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        phase_start_x + Inches(1.35),
        phase_top + Inches(0.55),
        Inches(11.0),
        Inches(0.06),
    )
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = LIGHT_GRAY
    line_shape.line.fill.background()

    for i, (num, title, color) in enumerate(phases):
        left = phase_start_x + i * (phase_width + phase_gap)

        # Number circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, left + Inches(0.9), phase_top, Inches(0.9), Inches(0.9)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()

        # Number text in circle
        tf = circle.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = num
        run.font.size = Pt(24)
        run.font.color.rgb = WHITE
        run.font.bold = True

        # Phase card below
        add_card(slide, left, phase_top + Inches(1.2), phase_width, Inches(2.0), accent_color=color, accent_position="top")

        # Phase title
        t_box = slide.shapes.add_textbox(left + Inches(0.2), phase_top + Inches(1.5), phase_width - Inches(0.4), Inches(1.5))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.size = Pt(14)
        run.font.color.rgb = WHITE
        run.font.bold = True

    add_footer(slide)

    # =========================================================================
    # SLIDE 10 - Closing
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_background_image(slide, "09_fresh.jpg")
    add_dark_overlay(slide, "medium")

    # Main title
    title_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(1.2))
    set_text_props(title_box.text_frame, "Smart Food Management", 48, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Tagline
    tag_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.5), Inches(10.333), Inches(0.8))
    set_text_props(tag_box.text_frame, "Every Meal Matters. Zero Waste. Full Plates.", 28, GREEN, bold=True, alignment=PP_ALIGN.CENTER)

    # JTP branding
    brand_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(10.333), Inches(0.6))
    set_text_props(brand_box.text_frame, "JTP Co., Ltd.", 24, WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_footer(slide)

    # =========================================================================
    # Save
    # =========================================================================
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SmartFoodManagement.pptx")
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    create_presentation()
